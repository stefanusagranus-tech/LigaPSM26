"""
ppt_generator.py
================
Modul generator laporan PPT untuk LigaPSM.
Dipisah dari main file biar gak makan space.
"""

import io
import streamlit as st

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt


def generate_ppt_report(
    title, month_year_str, generated_time_str,
    psm_data, pwp_data, sueger_data, sg_data, ceban_data,
    total_poin, top3_kasir,
):
    """
    Generate PPT Report Pro dengan chart & visual (9 slide).
    """
    try:
        # =========================================================
        # 🎨 WARNA
        # =========================================================
        CLR_DARK_NAVY = RGBColor(0x0C, 0x14, 0x27)
        CLR_NAVY = RGBColor(0x16, 0x24, 0x47)
        CLR_GOLD = RGBColor(0xD4, 0xAF, 0x37)
        CLR_GOLD_LIGHT = RGBColor(0xF7, 0xE7, 0xB4)
        CLR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
        CLR_TEXT = RGBColor(0xF1, 0xE5, 0xC7)
        CLR_GREEN = RGBColor(0x10, 0xB9, 0x81)
        CLR_BLUE = RGBColor(0x3B, 0x82, 0xF6)
        CLR_PURPLE = RGBColor(0xA8, 0x55, 0xF7)
        CLR_ORANGE = RGBColor(0xF5, 0x9E, 0x0B)
        CLR_PINK = RGBColor(0xEC, 0x48, 0x99)
        CLR_GRAY = RGBColor(0x94, 0xA3, 0xB8)

        # =========================================================
        # 📊 SETUP PRESENTATION
        # =========================================================
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        BLANK = prs.slide_layouts[6]

        # =========================================================
        # 🛠️ HELPER FUNCTIONS
        # =========================================================
        def add_bg(slide, color=CLR_DARK_NAVY):
            bg = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
            )
            bg.fill.solid()
            bg.fill.fore_color.rgb = color
            bg.line.fill.background()
            return bg

        def add_text(slide, text, left, top, width, height,
                     font_size=18, bold=False, color=CLR_TEXT,
                     align=PP_ALIGN.LEFT, font_name="Georgia"):
            tb = slide.shapes.add_textbox(left, top, width, height)
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = 0
            tf.margin_right = 0
            p = tf.paragraphs[0]
            p.alignment = align
            r = p.add_run()
            r.text = str(text)
            r.font.name = font_name
            r.font.size = Pt(font_size)
            r.font.bold = bold
            r.font.color.rgb = color
            return tb

        def add_progress_bar(slide, left, top, width, height,
                             percent, color_fill=CLR_GREEN):
            bg = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
            )
            bg.fill.solid()
            bg.fill.fore_color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            bg.line.color.rgb = CLR_GOLD
            bg.line.width = Pt(1)

            percent_clamped = max(0, min(100, percent))
            fill_width = int(width * (percent_clamped / 100))
            if fill_width > 0:
                fill = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    left, top, fill_width, height
                )
                fill.fill.solid()
                fill.fill.fore_color.rgb = color_fill
                fill.line.fill.background()

        def add_gold_border(slide):
            border = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(0.15), Inches(0.15),
                prs.slide_width - Inches(0.3),
                prs.slide_height - Inches(0.3),
            )
            border.fill.background()
            border.line.color.rgb = CLR_GOLD
            border.line.width = Pt(2.5)

        def add_chart_image(slide, fig, left, top, width, height):
            buf = io.BytesIO()
            fig.savefig(buf, format="png", dpi=150,
                        bbox_inches="tight", transparent=True)
            buf.seek(0)
            plt.close(fig)
            slide.shapes.add_picture(buf, left, top, width, height)

        # =========================================================
        # SLIDE 1: COVER
        # =========================================================
        slide = prs.slides.add_slide(BLANK)
        add_bg(slide)
        add_gold_border(slide)

        add_text(slide, "REPORT SUMMARY PENJUALAN",
                 Inches(0.5), Inches(2.1), Inches(12.33), Inches(0.9),
                 font_size=36, bold=True, color=CLR_GOLD_LIGHT,
                 align=PP_ALIGN.CENTER)
        add_text(slide, month_year_str.upper(),
                 Inches(0.5), Inches(3.05), Inches(12.33), Inches(0.7),
                 font_size=28, bold=True, color=CLR_GOLD,
                 align=PP_ALIGN.CENTER)
        add_text(slide, "TOKO C383 - KARANG SATRIA",
                 Inches(0.5), Inches(3.9), Inches(12.33), Inches(0.6),
                 font_size=20, color=CLR_WHITE,
                 align=PP_ALIGN.CENTER)

        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(4.5), Inches(4.7), Inches(4.33), Pt(2)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = CLR_GOLD
        line.line.fill.background()

        add_text(slide, f"TOTAL POIN: {total_poin:.2f}",
                 Inches(0.5), Inches(5.0), Inches(12.33), Inches(0.9),
                 font_size=32, bold=True, color=CLR_GOLD_LIGHT,
                 align=PP_ALIGN.CENTER)

        add_text(slide, f"Generated: {generated_time_str}",
                 Inches(0.5), Inches(6.5), Inches(12.33), Inches(0.4),
                 font_size=12, color=CLR_GRAY,
                 align=PP_ALIGN.CENTER, font_name="Consolas")

        # =========================================================
        # SLIDE 2: OVERVIEW CHART
        # =========================================================
        slide = prs.slides.add_slide(BLANK)
        add_bg(slide)
        add_gold_border(slide)

        add_text(slide, "OVERVIEW SEMUA PROGRAM",
                 Inches(0.5), Inches(0.4), Inches(12.33), Inches(0.7),
                 font_size=28, bold=True, color=CLR_GOLD_LIGHT,
                 align=PP_ALIGN.CENTER)

        labels = ["PSM", "PWP", "Sueger", "SG", "Ceban"]
        achievement = [
            psm_data["ach"], pwp_data["ach_qty"], sueger_data["ach"],
            sg_data["ach"], ceban_data["ach"]
        ]
        colors = ["#3B82F6", "#A855F7", "#10B981", "#F59E0B", "#EC4899"]

        fig, ax = plt.subplots(figsize=(10, 4.5), facecolor="#0C1427")
        ax.set_facecolor("#0C1427")
        bars = ax.bar(labels, achievement, color=colors,
                      edgecolor="#D4AF37", linewidth=1.5)
        ax.axhline(y=100, color="#D4AF37", linestyle="--",
                   linewidth=1.5, alpha=0.6, label="Target 100%")
        ax.set_ylabel("Achievement (%)", color="#F1E5C7", fontsize=12)
        ax.set_title("Achievement per Program",
                     color="#F7E7B4", fontsize=16, fontweight="bold")
        ax.tick_params(colors="#F1E5C7", labelsize=11)
        for spine in ax.spines.values():
            spine.set_color("#D4AF37")
            spine.set_linewidth(1)
        ax.grid(axis="y", color="#334155", linestyle=":", alpha=0.4)
        ax.legend(facecolor="#1E293B", edgecolor="#D4AF37",
                  labelcolor="#F1E5C7", fontsize=10)

        for bar, val in zip(bars, achievement):
            h = bar.get_height()
            ax.text(bar.get_x() + bar.get_width() / 2, h + 2,
                    f"{val:.1f}%", ha="center", va="bottom",
                    color="#F7E7B4", fontsize=10, fontweight="bold")

        add_chart_image(slide, fig,
                        Inches(1.5), Inches(1.3),
                        Inches(10.33), Inches(5.5))

        # =========================================================
        # SLIDE 3-7: DETAIL PER PROGRAM
        # =========================================================
        def add_program_slide(title_text, main_data, extra_lines,
                              color_accent=CLR_GREEN, page_num=3):
            slide = prs.slides.add_slide(BLANK)
            add_bg(slide)
            add_gold_border(slide)

            add_text(slide, title_text,
                     Inches(0.6), Inches(0.5), Inches(12.13), Inches(0.8),
                     font_size=28, bold=True, color=CLR_GOLD_LIGHT,
                     align=PP_ALIGN.CENTER)

            line = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(1.5), Inches(1.35), Inches(10.33), Pt(2)
            )
            line.fill.solid()
            line.fill.fore_color.rgb = CLR_GOLD
            line.line.fill.background()

            add_text(slide, "ACHIEVEMENT",
                     Inches(0.8), Inches(1.7), Inches(4), Inches(0.5),
                     font_size=16, bold=True, color=CLR_GOLD,
                     font_name="Consolas")
            add_text(slide, f"{main_data['ach']:.1f}%",
                     Inches(9.5), Inches(1.6), Inches(3), Inches(0.7),
                     font_size=32, bold=True, color=color_accent,
                     align=PP_ALIGN.RIGHT, font_name="Consolas")

            add_progress_bar(
                slide,
                Inches(0.8), Inches(2.5),
                Inches(11.7), Inches(0.55),
                main_data["ach"],
                color_fill=color_accent
            )

            y_start = 3.5
            stat_items = main_data.get("stats", [])
            col1_x, col2_x = Inches(0.8), Inches(6.8)

            for i, (label, value) in enumerate(stat_items):
                x = col1_x if i % 2 == 0 else col2_x
                y = y_start + (i // 2) * 0.75

                box = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    x, Inches(y), Inches(5.7), Inches(0.65)
                )
                box.fill.solid()
                box.fill.fore_color.rgb = CLR_NAVY
                box.line.color.rgb = CLR_GOLD
                box.line.width = Pt(1)

                add_text(slide, label,
                         x + Inches(0.2), Inches(y + 0.05),
                         Inches(3), Inches(0.5),
                         font_size=14, color=CLR_TEXT,
                         font_name="Consolas")
                add_text(slide, value,
                         x + Inches(3.3), Inches(y + 0.05),
                         Inches(2.2), Inches(0.5),
                         font_size=16, bold=True, color=CLR_GOLD_LIGHT,
                         align=PP_ALIGN.RIGHT, font_name="Consolas")

            if extra_lines:
                add_text(slide, "TARGET HARIAN & SHIFT",
                         Inches(0.8), Inches(5.7), Inches(11.7), Inches(0.4),
                         font_size=14, bold=True, color=CLR_GOLD,
                         font_name="Consolas")
                add_text(slide, extra_lines,
                         Inches(0.8), Inches(6.1), Inches(11.7), Inches(1.2),
                         font_size=12, color=CLR_TEXT,
                         font_name="Consolas")

            add_text(slide, f"- Halaman {page_num} -",
                     Inches(0.5), Inches(7.1), Inches(12.33), Inches(0.3),
                     font_size=10, color=CLR_GRAY,
                     align=PP_ALIGN.CENTER, font_name="Consolas")

        psm_extra = (
            f"Target Harian: {psm_data['harian']} Pcs/hari\n"
            f"Shift 1 (40%): {psm_data['shift1']} Pcs   |   "
            f"Shift 2 (40%): {psm_data['shift2']} Pcs   |   "
            f"Shift 3 (20%): {psm_data['shift3']} Pcs"
        )
        add_program_slide(
            "PROGRAM PSM",
            {
                "ach": psm_data["ach"],
                "stats": [
                    ("Target PSM", f"{psm_data['target']} Pcs"),
                    ("Actual Qty", f"{psm_data['actual']} Pcs"),
                    ("Achievement", f"{psm_data['ach']:.1f}%"),
                    ("Poin PSM", f"{psm_data['poin']:.2f} / 20"),
                ]
            },
            psm_extra, color_accent=CLR_BLUE, page_num=3
        )

        pwp_extra = (
            f"Target Harian: {pwp_data['harian']} Pcs/hari\n"
            f"Shift 1 (40%): {pwp_data['shift1']} Pcs   |   "
            f"Shift 2 (40%): {pwp_data['shift2']} Pcs   |   "
            f"Shift 3 (20%): {pwp_data['shift3']} Pcs"
        )
        add_program_slide(
            "PROGRAM PWP",
            {
                "ach": pwp_data["ach_qty"],
                "stats": [
                    ("Syarat Redeem", f"{pwp_data['syarat']}"),
                    ("Total Redeem", f"{pwp_data['redeem']}"),
                    ("Target Qty", f"{pwp_data['target']} Pcs"),
                    ("Actual Qty", f"{pwp_data['actual']} Pcs"),
                    ("Ach. Redeem", f"{pwp_data['ach_redeem']:.1f}%"),
                    ("Poin PWP", f"{pwp_data['poin']:.2f} / 25"),
                ]
            },
            pwp_extra, color_accent=CLR_PURPLE, page_num=4
        )

        add_program_slide(
            "PROGRAM SUEGER",
            {
                "ach": sueger_data["ach"],
                "stats": [
                    ("Syarat Redeem", f"{sueger_data['syarat']}"),
                    ("Qty Redeem", f"{sueger_data['redeem']}"),
                    ("Achievement", f"{sueger_data['ach']:.1f}%"),
                ]
            },
            None, color_accent=CLR_GREEN, page_num=5
        )

        sg_extra = (
            f"Target Harian: {sg_data['harian']} Pcs/hari\n"
            f"Shift 1 (40%): {sg_data['shift1']} Pcs   |   "
            f"Shift 2 (40%): {sg_data['shift2']} Pcs   |   "
            f"Shift 3 (20%): {sg_data['shift3']} Pcs"
        )
        add_program_slide(
            "PROGRAM SERBA GRATIS (SG)",
            {
                "ach": sg_data["ach"],
                "stats": [
                    ("Target Qty", f"{sg_data['target']} Pcs"),
                    ("Actual Qty", f"{sg_data['actual']} Pcs"),
                    ("Achievement", f"{sg_data['ach']:.1f}%"),
                    ("Poin SG", f"{sg_data['poin']:.2f} / 30"),
                ]
            },
            sg_extra, color_accent=CLR_ORANGE, page_num=6
        )

        add_program_slide(
            "CEMILAN CEBAN",
            {
                "ach": ceban_data["ach"],
                "stats": [
                    ("Target Qty", f"{ceban_data['target']} Pcs"),
                    ("Actual Qty", f"{ceban_data['actual']} Pcs"),
                    ("Achievement", f"{ceban_data['ach']:.1f}%"),
                ]
            },
            None, color_accent=CLR_PINK, page_num=7
        )

        # =========================================================
        # SLIDE 8: LEADERBOARD
        # =========================================================
        slide = prs.slides.add_slide(BLANK)
        add_bg(slide)
        add_gold_border(slide)

        add_text(slide, "LEADERBOARD TOP 3",
                 Inches(0.5), Inches(0.4), Inches(12.33), Inches(0.8),
                 font_size=28, bold=True, color=CLR_GOLD_LIGHT,
                 align=PP_ALIGN.CENTER)

        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(1.5), Inches(1.35), Inches(10.33), Pt(2)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = CLR_GOLD
        line.line.fill.background()

        podium_positions = [
            (2, 0.8, 3.3, 2.3, RGBColor(0xC0, 0xC0, 0xC0)),
            (1, 4.7, 2.3, 3.3, CLR_GOLD),
            (3, 8.6, 3.8, 1.8, RGBColor(0xCD, 0x7F, 0x32)),
        ]

        for rank, x, y, h, color in podium_positions:
            block = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x), Inches(y),
                Inches(4.0), Inches(h)
            )
            block.fill.solid()
            block.fill.fore_color.rgb = CLR_NAVY
            block.line.color.rgb = color
            block.line.width = Pt(3)

            if rank <= len(top3_kasir):
                nama, nilai, satuan = top3_kasir[rank - 1]
            else:
                nama, nilai, satuan = "- Belum ada data -", 0, ""

            add_text(slide, f"#{rank}",
                     Inches(x), Inches(y + 0.15),
                     Inches(4.0), Inches(0.5),
                     font_size=18, bold=True, color=color,
                     align=PP_ALIGN.CENTER, font_name="Consolas")

            add_text(slide, str(nama),
                     Inches(x), Inches(y + 0.75),
                     Inches(4.0), Inches(0.7),
                     font_size=18, bold=True, color=CLR_WHITE,
                     align=PP_ALIGN.CENTER, font_name="Consolas")

            add_text(slide, f"{nilai} {satuan}",
                     Inches(x), Inches(y + 1.5),
                     Inches(4.0), Inches(0.6),
                     font_size=22, bold=True, color=CLR_GOLD_LIGHT,
                     align=PP_ALIGN.CENTER, font_name="Consolas")

        # =========================================================
        # SLIDE 9: INSIGHT
        # =========================================================
        slide = prs.slides.add_slide(BLANK)
        add_bg(slide)
        add_gold_border(slide)

        add_text(slide, "REKOMENDASI & INSIGHT",
                 Inches(0.5), Inches(0.4), Inches(12.33), Inches(0.8),
                 font_size=28, bold=True, color=CLR_GOLD_LIGHT,
                 align=PP_ALIGN.CENTER)

        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(1.5), Inches(1.35), Inches(10.33), Pt(2)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = CLR_GOLD
        line.line.fill.background()

        all_programs = [
            ("PSM", psm_data["ach"]),
            ("PWP", pwp_data["ach_qty"]),
            ("Sueger", sueger_data["ach"]),
            ("SG", sg_data["ach"]),
            ("Ceban", ceban_data["ach"]),
        ]
        best = max(all_programs, key=lambda x: x[1])
        worst = min(all_programs, key=lambda x: x[1])
        avg_ach = sum(a for _, a in all_programs) / len(all_programs)

        insight_list = [
            f"[OK] Program terbaik: {best[0]} dengan achievement {best[1]:.1f}%",
            f"[!] Program butuh perhatian: {worst[0]} ({worst[1]:.1f}%)",
            f"Rata-rata achievement semua program: {avg_ach:.1f}%",
            "",
            "REKOMENDASI ACTION:",
        ]
        if worst[1] < 50:
            insight_list.append(
                f"  - Fokus extra untuk program {worst[0]} - review strategi"
            )
        if avg_ach < 70:
            insight_list.append(
                "  - Adakan briefing mingguan untuk evaluasi progres"
            )
        if best[1] >= 90:
            insight_list.append(
                f"  - Program {best[0]} sudah excellent - jadikan benchmark"
            )
        insight_list.append("")
        insight_list.append(f"TOTAL POIN KESELURUHAN: {total_poin:.2f}")

        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(1.0), Inches(1.8),
            Inches(11.33), Inches(5.3)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = CLR_NAVY
        box.line.color.rgb = CLR_GOLD
        box.line.width = Pt(2)

        tb = slide.shapes.add_textbox(
            Inches(1.4), Inches(2.1), Inches(10.53), Inches(4.8)
        )
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line_txt in enumerate(insight_list):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = line_txt
            r.font.name = "Consolas"
            r.font.size = Pt(14)
            r.font.color.rgb = CLR_TEXT

        # === OUTPUT ===
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        return buf.getvalue()

    except Exception as e:
        st.error(f"Gagal generate PPT: {e}")
        import traceback
        st.code(traceback.format_exc())
        return None
